"""Generate an Excel review workbook for the OPRO CronJob Monitor app."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Style constants ──
HEADER_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4A6741", end_color="4A6741", fill_type="solid")  # Ashley green
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="3E2723")
SUBTITLE_FONT = Font(name="Calibri", bold=True, size=11, color="5D4037")
NORMAL_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
ALT_FILL = PatternFill(start_color="F5F5F0", end_color="F5F5F0", fill_type="solid")


def style_header_row(ws, cols, row=1):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def auto_width(ws, min_w=14, max_w=40):
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        best = min_w
        for cell in col_cells:
            if cell.value:
                best = max(best, min(len(str(cell.value)) + 4, max_w))
        ws.column_dimensions[col_letter].width = best


# ═══════════════════════════════════════════
# Sheet 1 – Overview
# ═══════════════════════════════════════════
ws1 = wb.active
ws1.title = "Overview"

overview_data = [
    ("OPRO CronJob Monitor – Application Overview", ""),
    ("", ""),
    ("Item", "Details"),
    ("Application Name", "OPRO CronJob Monitor"),
    ("Purpose", "Monitor the execution of OPRO batch cronjobs for Ashley Furniture Industries. "
                "Each batch runs a six-step pipeline that processes furniture orders, allocates inventory, "
                "and generates CODIS suggestions."),
    ("Tech Stack", "React 19 + TypeScript · Vite · Material UI (MUI) · React Query · React Router · date-fns"),
    ("Theme", "Ashley Furniture branded – Brown / Green MUI color palette"),
    ("", ""),
    ("Screen", "Description"),
    ("Screen 1 – Batch List Page", "Paginated table of all batch executions with filtering by Job ID, date, "
                                    "time frame, and status. URL query-param persistence. Displays batch count, "
                                    "status chips, duration, and order counts."),
    ("Screen 2 – Batch Details Page", "Full analytics for a single batch: Batch header with status & timing, "
                                       "Metric cards (Orders Processed, Allocated, CODIS Suggestions), "
                                       "Batch summary table, Step-by-step execution timeline with progress bars. "
                                       "Actions: Export PDF, Refresh, Email Report, Download CSV."),
    ("", ""),
    ("Key Features", ""),
    ("", "• Paginated & filterable batch list"),
    ("", "• URL-persisted filters"),
    ("", "• Batch detail drill-down with metric cards"),
    ("", "• Execution timeline with step progress bars"),
    ("", "• CSV download of batch logs"),
    ("", "• Loading skeletons & error states"),
    ("", "• Accessibility (ARIA labels, keyboard nav, WCAG AA)"),
    ("", "• Mock data (20 batches) for demo"),
]

for r, (a, b) in enumerate(overview_data, 1):
    ws1.cell(row=r, column=1, value=a).font = SUBTITLE_FONT if r in (3, 9, 13) else (TITLE_FONT if r == 1 else NORMAL_FONT)
    ws1.cell(row=r, column=2, value=b).font = NORMAL_FONT
    ws1.cell(row=r, column=1).alignment = WRAP
    ws1.cell(row=r, column=2).alignment = WRAP

ws1.column_dimensions["A"].width = 35
ws1.column_dimensions["B"].width = 90

# ═══════════════════════════════════════════
# Sheet 2 – Action Items
# ═══════════════════════════════════════════
ws2 = wb.create_sheet("Action Items")
headers2 = ["Action Item", "Status"]
for c, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=c, value=h)
style_header_row(ws2, len(headers2))

# Pre-fill a few placeholder rows so the sheet isn't empty
placeholders = [
    ("", "Not Started"),
    ("", "Not Started"),
    ("", "Not Started"),
]
for r, (item, status) in enumerate(placeholders, 2):
    ws2.cell(row=r, column=1, value=item).font = NORMAL_FONT
    ws2.cell(row=r, column=2, value=status).font = NORMAL_FONT
    ws2.cell(row=r, column=1).alignment = WRAP
    ws2.cell(row=r, column=2).alignment = CENTER
    for c in (1, 2):
        ws2.cell(row=r, column=c).border = THIN_BORDER

ws2.column_dimensions["A"].width = 60
ws2.column_dimensions["B"].width = 20

# ═══════════════════════════════════════════
# Sheet 3 & 4 – Screen review sheets
# ═══════════════════════════════════════════
screen_headers = ["Date", "Screenshot", "Description", "Comments", "Commented By", "Dev Feedback"]

for sheet_name in ("BatchListPage", "BatchDetailsPage"):
    ws = wb.create_sheet(sheet_name)
    for c, h in enumerate(screen_headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header_row(ws, len(screen_headers))

    # Add 10 empty rows with borders so the team can fill in
    for r in range(2, 12):
        for c in range(1, len(screen_headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = NORMAL_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if r % 2 == 0:
                cell.fill = ALT_FILL

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 40
    ws.row_dimensions[1].height = 28
    for r in range(2, 12):
        ws.row_dimensions[r].height = 80  # tall rows for screenshots

OUT = "OPRO_CronJob_Monitor_Review.xlsx"
wb.save(OUT)
print(f"✅ Created {OUT}")
